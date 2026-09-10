# ============================================================
# PRODUCTION PLANT CARBON FOOTPRINT CHALLENGE
# ROUND 3 - CLIMATE ACTION SIMULATION
# ============================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from tabulate import tabulate
from sortedcontainers import SortedDict
from openpyxl import load_workbook
import pyarrow as pa


# ============================================================
# 1. EMISSION FACTORS
# ============================================================

ELECTRICITY_FACTOR = 0.50
FUEL_FACTOR = 2.68
RAW_TRANSPORT_FACTOR = 0.12
PRODUCT_TRANSPORT_FACTOR = 0.10
WASTE_FACTOR = 0.80
PACKAGING_FACTOR = 1.50


# ============================================================
# 2. REUSABLE EMISSION FUNCTION
# ============================================================

def calculate_emissions(record):
    """
    Calculate all emission sources for one production record.
    """

    electricity = (
        record["electricity"]
        * ELECTRICITY_FACTOR
    )

    fuel = (
        record["fuel"]
        * FUEL_FACTOR
    )

    raw_transport = (
        record["raw_transport"]
        * RAW_TRANSPORT_FACTOR
    )

    product_transport = (
        record["product_transport"]
        * PRODUCT_TRANSPORT_FACTOR
    )

    waste = (
        record["waste"]
        * WASTE_FACTOR
    )

    packaging = (
        record["packaging"]
        * PACKAGING_FACTOR
    )

    return {
        "Electricity": electricity,
        "Fuel": fuel,
        "Raw Material Transportation": raw_transport,
        "Product Transportation": product_transport,
        "Waste": waste,
        "Packaging": packaging
    }


# ============================================================
# 3. LOAD DATASET
# ============================================================

records = [
    {
        "batch_id": "B001",
        "line": "Line A",
        "units": 1000,
        "electricity": 2500,
        "fuel": 400,
        "raw_transport": 600,
        "product_transport": 500,
        "waste": 120,
        "packaging": 300
    },

    {
        "batch_id": "B002",
        "line": "Line B",
        "units": 1200,
        "electricity": 3000,
        "fuel": 450,
        "raw_transport": 700,
        "product_transport": 650,
        "waste": 150,
        "packaging": 350
    },

    {
        "batch_id": "B003",
        "line": "Line C",
        "units": 900,
        "electricity": 2100,
        "fuel": 350,
        "raw_transport": 500,
        "product_transport": 450,
        "waste": 90,
        "packaging": 280
    },

    {
        "batch_id": "B004",
        "line": "Line A",
        "units": 1500,
        "electricity": 3800,
        "fuel": 520,
        "raw_transport": 850,
        "product_transport": 800,
        "waste": 210,
        "packaging": 420
    },

    {
        "batch_id": "B005",
        "line": "Line B",
        "units": 1100,
        "electricity": 2700,
        "fuel": 390,
        "raw_transport": 620,
        "product_transport": 550,
        "waste": 130,
        "packaging": 310
    }
]


# ============================================================
# 4. PROCESS EVERY RECORD
# ============================================================

results = []

for record in records:

    emissions = calculate_emissions(record)

    total_carbon = sum(emissions.values())

    carbon_per_unit = (
        total_carbon
        / record["units"]
    )

    results.append({
        "Batch ID": record["batch_id"],
        "Production Line": record["line"],
        "Units": record["units"],
        "Electricity": emissions["Electricity"],
        "Fuel": emissions["Fuel"],
        "Raw Material Transportation":
            emissions["Raw Material Transportation"],
        "Product Transportation":
            emissions["Product Transportation"],
        "Waste":
            emissions["Waste"],
        "Packaging":
            emissions["Packaging"],
        "Total Carbon":
            total_carbon,
        "Carbon Per Unit":
            carbon_per_unit,
        "Waste (kg)":
            record["waste"]
    })


# ============================================================
# 5. CREATE DATAFRAME
# ============================================================

df = pd.DataFrame(results)


# ============================================================
# 6. TOTAL PLANT CARBON FOOTPRINT
# ============================================================

total_carbon = df["Total Carbon"].sum()


# ============================================================
# 7. AVERAGE CARBON PER BATCH
# ============================================================

average_carbon = df["Total Carbon"].mean()


# ============================================================
# 8. TOTAL UNITS
# ============================================================

total_units = df["Units"].sum()


# ============================================================
# 9. CARBON PER UNIT
# ============================================================

plant_carbon_per_unit = (
    total_carbon
    / total_units
)


# ============================================================
# 10. HIGHEST-EMISSION BATCH
# ============================================================

highest_batch = df.loc[
    df["Total Carbon"].idxmax()
]


# ============================================================
# 11. LOWEST-EMISSION BATCH
# ============================================================

lowest_batch = df.loc[
    df["Total Carbon"].idxmin()
]


# ============================================================
# 12. MOST EFFICIENT BATCH
# ============================================================

efficient_batch = df.loc[
    df["Carbon Per Unit"].idxmin()
]


# ============================================================
# 13. PRODUCTION-LINE COMPARISON
# ============================================================

line_totals = (
    df.groupby("Production Line")["Total Carbon"]
    .sum()
)

highest_line = line_totals.idxmax()


# ============================================================
# 14. EMISSION SOURCE ANALYSIS
# ============================================================

source_columns = [
    "Electricity",
    "Fuel",
    "Raw Material Transportation",
    "Product Transportation",
    "Waste",
    "Packaging"
]

source_totals = df[source_columns].sum()

largest_source = source_totals.idxmax()


# ============================================================
# 15. WASTE ANALYSIS
# ============================================================

total_waste = df["Waste (kg)"].sum()

average_waste = df["Waste (kg)"].mean()

highest_waste_batch = df.loc[
    df["Waste (kg)"].idxmax()
]


# ============================================================
# 16. SUSTAINABLE PRODUCTION INDICATOR
# ============================================================

carbon_efficiency = (
    total_units
    / total_carbon
)


# ============================================================
# 17. CLIMATE ACTION SIMULATION
# ============================================================

# Improvements:
# Electricity       -> 15% reduction
# Fuel              -> 10% reduction
# Waste             -> 20% reduction
# Product Transport -> 10% reduction


projected_electricity = (
    df["Electricity"].sum()
    * 0.85
)

projected_fuel = (
    df["Fuel"].sum()
    * 0.90
)

projected_raw_transport = (
    df["Raw Material Transportation"].sum()
)

projected_product_transport = (
    df["Product Transportation"].sum()
    * 0.90
)

projected_waste = (
    df["Waste"].sum()
    * 0.80
)

projected_packaging = (
    df["Packaging"].sum()
)


# ============================================================
# 18. PROJECTED CARBON FOOTPRINT
# ============================================================

projected_carbon = (
    projected_electricity
    + projected_fuel
    + projected_raw_transport
    + projected_product_transport
    + projected_waste
    + projected_packaging
)


# ============================================================
# 19. CO2 REDUCTION
# ============================================================

carbon_reduction = (
    total_carbon
    - projected_carbon
)


# ============================================================
# 20. PERCENTAGE REDUCTION
# ============================================================

percentage_reduction = (
    carbon_reduction
    / total_carbon
) * 100


# ============================================================
# 21. TERMINAL REPORT
# ============================================================

print("\n" + "=" * 80)
print("       PRODUCTION PLANT CARBON FOOTPRINT")
print("       ROUND 3 - CLIMATE ACTION SIMULATION")
print("=" * 80)


# ------------------------------------------------------------
# Plant Overview
# ------------------------------------------------------------

print("\nPLANT OVERVIEW")
print("-" * 80)

overview = [
    ["Total Batches", len(df)],
    ["Total Units Produced", f"{total_units:,}"],
    ["Total Carbon Footprint",
     f"{total_carbon:,.2f} kg CO₂"],
    ["Average Carbon per Batch",
     f"{average_carbon:,.2f} kg CO₂"],
    ["Carbon per Unit",
     f"{plant_carbon_per_unit:.4f} kg CO₂"],
    ["Carbon Efficiency",
     f"{carbon_efficiency:.4f} units/kg CO₂"]
]

print(
    tabulate(
        overview,
        headers=["Indicator", "Result"],
        tablefmt="grid"
    )
)


# ------------------------------------------------------------
# Batch Analysis
# ------------------------------------------------------------

print("\nBATCH ANALYSIS")
print("-" * 80)

batch_table = df[
    [
        "Batch ID",
        "Production Line",
        "Units",
        "Total Carbon",
        "Carbon Per Unit"
    ]
].copy()

batch_table["Total Carbon"] = (
    batch_table["Total Carbon"]
    .round(2)
)

batch_table["Carbon Per Unit"] = (
    batch_table["Carbon Per Unit"]
    .round(4)
)

print(
    tabulate(
        batch_table,
        headers="keys",
        tablefmt="grid",
        showindex=False
    )
)


# ------------------------------------------------------------
# Highest / Lowest / Efficient
# ------------------------------------------------------------

print("\nBATCH RESULTS")
print("-" * 80)

print(
    f"Highest-Emission Batch : "
    f"{highest_batch['Batch ID']}"
)

print(
    f"Production Line        : "
    f"{highest_batch['Production Line']}"
)

print(
    f"Carbon Footprint       : "
    f"{highest_batch['Total Carbon']:.2f} kg CO₂"
)

print(
    f"\nLowest-Emission Batch  : "
    f"{lowest_batch['Batch ID']}"
)

print(
    f"Carbon Footprint       : "
    f"{lowest_batch['Total Carbon']:.2f} kg CO₂"
)

print(
    f"\nMost Efficient Batch   : "
    f"{efficient_batch['Batch ID']}"
)

print(
    f"Carbon Per Unit        : "
    f"{efficient_batch['Carbon Per Unit']:.4f} kg CO₂"
)


# ------------------------------------------------------------
# Production Line Analysis
# ------------------------------------------------------------

print("\nPRODUCTION-LINE COMPARISON")
print("-" * 80)

line_table = []

for line, value in line_totals.items():

    line_table.append([
        line,
        f"{value:,.2f} kg CO₂"
    ])

print(
    tabulate(
        line_table,
        headers=["Production Line", "Total Carbon"],
        tablefmt="grid"
    )
)

print(
    f"\nHighest-Emission Line: {highest_line}"
)


# ------------------------------------------------------------
# Emission Source Analysis
# ------------------------------------------------------------

print("\nEMISSION SOURCE ANALYSIS")
print("-" * 80)

source_table = []

for source, value in source_totals.items():

    source_table.append([
        source,
        f"{value:,.2f} kg CO₂",
        f"{(value / total_carbon) * 100:.2f}%"
    ])

print(
    tabulate(
        source_table,
        headers=[
            "Emission Source",
            "Total Emissions",
            "% of Total"
        ],
        tablefmt="grid"
    )
)

print(
    f"\nLargest Contributor: {largest_source}"
)


# ------------------------------------------------------------
# Waste Analysis
# ------------------------------------------------------------

print("\nWASTE ANALYSIS")
print("-" * 80)

print(
    f"Total Waste          : "
    f"{total_waste:,.2f} kg"
)

print(
    f"Average Waste/Batch  : "
    f"{average_waste:,.2f} kg"
)

print(
    f"Highest Waste Batch  : "
    f"{highest_waste_batch['Batch ID']}"
)

print(
    f"Waste Produced       : "
    f"{highest_waste_batch['Waste (kg)']:.2f} kg"
)


# ============================================================
# 22. CLIMATE ACTION RESULTS
# ============================================================

print("\n" + "=" * 80)
print("       CLIMATE ACTION SIMULATION")
print("=" * 80)

simulation_table = [
    [
        "Original Footprint",
        f"{total_carbon:,.2f} kg CO₂"
    ],

    [
        "Projected Footprint",
        f"{projected_carbon:,.2f} kg CO₂"
    ],

    [
        "CO₂ Reduction",
        f"{carbon_reduction:,.2f} kg CO₂"
    ],

    [
        "Percentage Reduction",
        f"{percentage_reduction:.2f}%"
    ]
]

print(
    tabulate(
        simulation_table,
        headers=["Indicator", "Result"],
        tablefmt="grid"
    )
)


# ============================================================
# 23. PYARROW
# ============================================================

arrow_table = pa.Table.from_pandas(df)

print("\nPyArrow Dataset")
print(
    f"Rows: {arrow_table.num_rows}"
)

print(
    f"Columns: {arrow_table.num_columns}"
)


# ============================================================
# 24. SORTED CONTAINER
# ============================================================

sorted_batches = SortedDict()

for _, row in df.iterrows():

    sorted_batches[
        row["Batch ID"]
    ] = row["Total Carbon"]


print("\nSORTED BATCH CARBON FOOTPRINT")

for batch, carbon in sorted_batches.items():

    print(
        f"{batch}: {carbon:,.2f} kg CO₂"
    )


# ============================================================
# 25. NUMPY ANALYSIS
# ============================================================

carbon_array = np.array(
    df["Total Carbon"]
)

print("\nNUMPY ANALYSIS")

print(
    f"Maximum : {np.max(carbon_array):,.2f} kg CO₂"
)

print(
    f"Minimum : {np.min(carbon_array):,.2f} kg CO₂"
)

print(
    f"Mean    : {np.mean(carbon_array):,.2f} kg CO₂"
)


# ============================================================
# 26. EXCEL REPORT
# ============================================================

excel_file = "round3_carbon_footprint.xlsx"

with pd.ExcelWriter(
    excel_file,
    engine="xlsxwriter"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Batch Analysis",
        index=False
    )

    line_totals.to_frame(
        "Total Carbon"
    ).to_excel(
        writer,
        sheet_name="Production Lines"
    )

    source_totals.to_frame(
        "Total Emissions"
    ).to_excel(
        writer,
        sheet_name="Emission Sources"
    )

    workbook = writer.book

    worksheet = writer.sheets[
        "Batch Analysis"
    ]

    header_format = workbook.add_format({
        "bold": True,
        "border": 1
    })

    worksheet.set_column(
        "A:A",
        12
    )

    worksheet.set_column(
        "B:B",
        18
    )

    worksheet.set_column(
        "C:K",
        20
    )

    # Summary section

    worksheet.write(
        "M2",
        "CLIMATE ACTION",
        header_format
    )

    worksheet.write(
        "M3",
        "Original Footprint"
    )

    worksheet.write(
        "N3",
        total_carbon
    )

    worksheet.write(
        "M4",
        "Projected Footprint"
    )

    worksheet.write(
        "N4",
        projected_carbon
    )

    worksheet.write(
        "M5",
        "CO₂ Reduction"
    )

    worksheet.write(
        "N5",
        carbon_reduction
    )

    worksheet.write(
        "M6",
        "Reduction %"
    )

    worksheet.write(
        "N6",
        percentage_reduction / 100
    )


# ============================================================
# 27. VERIFY EXCEL
# ============================================================

workbook = load_workbook(
    excel_file
)

print("\nExcel Report Created")
print(
    f"Sheets: {workbook.sheetnames}"
)


# ============================================================
# 28. VISUALIZATION
# ============================================================

plt.figure(figsize=(10, 6))

plt.bar(
    df["Batch ID"],
    df["Total Carbon"]
)

plt.title(
    "Carbon Footprint by Production Batch"
)

plt.xlabel(
    "Batch ID"
)

plt.ylabel(
    "Carbon Footprint (kg CO₂)"
)

plt.tight_layout()

plt.savefig(
    "round3_batch_carbon.png",
    dpi=150
)

plt.show()


# ============================================================
# FINAL
# ============================================================

print("\n" + "=" * 80)
print("ROUND 3 COMPLETE")
print("=" * 80)

print(
    f"Original Carbon Footprint : "
    f"{total_carbon:,.2f} kg CO₂"
)

print(
    f"Projected Carbon Footprint: "
    f"{projected_carbon:,.2f} kg CO₂"
)

print(
    f"Total CO₂ Reduction       : "
    f"{carbon_reduction:,.2f} kg CO₂"
)

print(
    f"Percentage Reduction      : "
    f"{percentage_reduction:.2f}%"
)
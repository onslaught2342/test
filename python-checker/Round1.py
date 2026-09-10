import sympy as sp
import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from tabulate import tabulate
from sortedcontainers import SortedDict
from openpyxl import load_workbook

import pyarrow as pa


# ============================================================
# PRODUCTION PLANT CARBON FOOTPRINT CHALLENGE
# ROUND 1 - SINGLE PRODUCTION BATCH
# ============================================================


# -----------------------------
# 1. INPUT DATA
# -----------------------------

units = 1000
electricity = 2500       # kWh
fuel = 400               # litres
raw_material = 1500      # kg
raw_transport = 300      # ton-km
product_transport = 500  # ton-km
waste = 120              # kg
packaging = 80            # kg


# -----------------------------
# 2. EMISSION FACTORS
# -----------------------------

ELECTRICITY_FACTOR = 0.50
FUEL_FACTOR = 2.68
RAW_TRANSPORT_FACTOR = 0.12
PRODUCT_TRANSPORT_FACTOR = 0.10
WASTE_FACTOR = 0.80
PACKAGING_FACTOR = 1.50


# -----------------------------
# 3. FUNCTIONS
# -----------------------------

def calculate_emissions(electricity, fuel, raw_transport,
                        product_transport, waste, packaging):

    emissions = {
        "Electricity": electricity * ELECTRICITY_FACTOR,
        "Fuel/Gas": fuel * FUEL_FACTOR,
        "Raw Material Transport": raw_transport * RAW_TRANSPORT_FACTOR,
        "Product Transport": product_transport * PRODUCT_TRANSPORT_FACTOR,
        "Waste": waste * WASTE_FACTOR,
        "Packaging": packaging * PACKAGING_FACTOR
    }

    return emissions


def calculate_status(carbon_per_unit):

    if carbon_per_unit < 5:
        return "Low"
    elif carbon_per_unit < 10:
        return "Moderate"
    else:
        return "High"


def calculate_waste_percentage(waste, raw_material):

    return (waste / raw_material) * 100


# -----------------------------
# 4. CALCULATE EMISSIONS
# -----------------------------

emissions = calculate_emissions(
    electricity,
    fuel,
    raw_transport,
    product_transport,
    waste,
    packaging
)

total_carbon = sum(emissions.values())

carbon_per_unit = total_carbon / units

waste_percentage = calculate_waste_percentage(
    waste,
    raw_material
)

highest_source = max(
    emissions,
    key=emissions.get
)

plant_status = calculate_status(carbon_per_unit)


# -----------------------------
# 5. NUMPY CALCULATION
# -----------------------------

emission_values = np.array(list(emissions.values()))

average_emission = np.mean(emission_values)

maximum_emission = np.max(emission_values)


# -----------------------------
# 6. SYMPY CALCULATION
# -----------------------------

# Symbolic calculation of carbon per unit

C = sp.Symbol("C")
T = sp.Symbol("T")
U = sp.Symbol("U")

carbon_formula = sp.Eq(C, T / U)

symbolic_result = carbon_formula.subs({
    T: total_carbon,
    U: units
})


# -----------------------------
# 7. SORTEDDICT
# -----------------------------

# Automatically keeps emission sources sorted alphabetically

sorted_emissions = SortedDict(emissions)


# -----------------------------
# 8. TERMINAL OUTPUT
# -----------------------------

print("\n" + "=" * 60)
print("       PRODUCTION PLANT CARBON FOOTPRINT")
print("=" * 60)

print(f"\nUnits Produced       : {units:,}")
print(f"Electricity Used     : {electricity:,} kWh")
print(f"Fuel Used            : {fuel:,} L")
print(f"Raw Material         : {raw_material:,} kg")
print(f"Waste                : {waste:,} kg")

print("\n" + "-" * 60)
print("EMISSION BREAKDOWN")
print("-" * 60)

table = []

for source, value in sorted_emissions.items():
    table.append([
        source,
        f"{value:.2f} kg CO₂"
    ])

print(
    tabulate(
        table,
        headers=["Emission Source", "Emissions"],
        tablefmt="grid"
    )
)

print("\n" + "-" * 60)
print("RESULTS")
print("-" * 60)

print(f"Total Carbon Footprint : {total_carbon:.2f} kg CO₂")
print(f"Carbon per Unit        : {carbon_per_unit:.2f} kg CO₂")
print(f"Waste Percentage       : {waste_percentage:.2f}%")
print(f"Highest Emission       : {highest_source}")
print(f"Plant Status           : {plant_status}")

print("\nSymPy Formula:")
print(carbon_formula)

print(f"SymPy Result: {symbolic_result}")

print(f"\nAverage Emission Source : {average_emission:.2f} kg CO₂")
print(f"Maximum Emission Source: {maximum_emission:.2f} kg CO₂")


# -----------------------------
# 9. PANDAS DATAFRAME
# -----------------------------

df = pd.DataFrame({
    "Emission Source": list(emissions.keys()),
    "Emissions (kg CO2)": list(emissions.values())
})

df["Percentage of Total"] = (
    df["Emissions (kg CO2)"] / total_carbon
) * 100

print("\nPANDAS DATAFRAME")
print("-" * 60)

print(df.to_string(index=False))


# -----------------------------
# 10. PYARROW
# -----------------------------

# Convert the Pandas DataFrame into an Arrow table

arrow_table = pa.Table.from_pandas(df)

print("\nPyArrow Table Created:")
print(f"Rows    : {arrow_table.num_rows}")
print(f"Columns : {arrow_table.num_columns}")


# -----------------------------
# 11. EXCEL EXPORT
# -----------------------------

excel_file = "carbon_footprint_round1.xlsx"

with pd.ExcelWriter(
    excel_file,
    engine="xlsxwriter"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Carbon Footprint",
        index=False
    )

    workbook = writer.book
    worksheet = writer.sheets["Carbon Footprint"]

    # Formatting
    header_format = workbook.add_format({
        "bold": True,
        "border": 1
    })

    number_format = workbook.add_format({
        "num_format": "0.00"
    })

    percentage_format = workbook.add_format({
        "num_format": "0.00%"
    })

    # Header formatting
    for col_num, value in enumerate(df.columns):
        worksheet.write(0, col_num, value, header_format)

    # Column widths
    worksheet.set_column("A:A", 28)
    worksheet.set_column("B:B", 20)
    worksheet.set_column("C:C", 20)

    # Add summary
    worksheet.write("E2", "Total Carbon Footprint", header_format)
    worksheet.write("F2", total_carbon)

    worksheet.write("E3", "Carbon per Unit", header_format)
    worksheet.write("F3", carbon_per_unit)

    worksheet.write("E4", "Waste Percentage", header_format)
    worksheet.write("F4", waste_percentage / 100, percentage_format)

    worksheet.write("E5", "Highest Emission", header_format)
    worksheet.write("F5", highest_source)

    worksheet.write("E6", "Plant Status", header_format)
    worksheet.write("F6", plant_status)


# -----------------------------
# 12. OPENPYXL
# -----------------------------

# Re-open the Excel file to verify that it was created correctly

workbook = load_workbook(excel_file)

print("\nExcel Workbook:")
print(f"File: {excel_file}")
print(f"Sheets: {workbook.sheetnames}")


# -----------------------------
# 13. MATPLOTLIB CHART
# -----------------------------

plt.figure(figsize=(10, 6))

plt.bar(
    df["Emission Source"],
    df["Emissions (kg CO2)"]
)

plt.title("Production Plant Carbon Footprint")
plt.xlabel("Emission Source")
plt.ylabel("Emissions (kg CO₂)")

plt.xticks(
    rotation=35,
    ha="right"
)

plt.tight_layout()

plt.savefig(
    "carbon_footprint_chart.png",
    dpi=150
)

plt.show()


# -----------------------------
# 14. FINAL RESULT
# -----------------------------

print("\n" + "=" * 60)
print("FINAL RESULT")
print("=" * 60)

print(f"""
Total Carbon Footprint : {total_carbon:.2f} kg CO₂
Carbon Footprint/Unit  : {carbon_per_unit:.2f} kg CO₂
Waste Percentage       : {waste_percentage:.2f}%
Highest Emission       : {highest_source}
Plant Status           : {plant_status}

Excel Report           : {excel_file}
Chart                  : carbon_footprint_chart.png
""")
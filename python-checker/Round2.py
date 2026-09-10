# ============================================================
# PRODUCTION PLANT CARBON FOOTPRINT CHALLENGE
# ROUND 2 - MULTIPLE PRODUCTION BATCHES
# ============================================================

import pandas as pd
import numpy as np
import matplotlib.pyplot as plt

from tabulate import tabulate
from sortedcontainers import SortedDict
from openpyxl import load_workbook
import pyarrow as pa


# ============================================================
# 1. EMISSION FACTORS
# ============================================================

ELECTRICITY_FACTOR = 0.50
FUEL_FACTOR = 2.68
RAW_TRANSPORT_FACTOR = 0.12
PRODUCT_TRANSPORT_FACTOR = 0.10
WASTE_FACTOR = 0.80
PACKAGING_FACTOR = 1.50


# ============================================================
# 2. PRODUCTION BATCH DATA
# ============================================================

batches = {
    "A": {
        "units": 1000,
        "electricity": 2500,
        "fuel": 400,
        "raw_transport": 600,
        "product_transport": 500,
        "waste": 120,
        "packaging": 300
    },

    "B": {
        "units": 1200,
        "electricity": 3000,
        "fuel": 450,
        "raw_transport": 700,
        "product_transport": 650,
        "waste": 150,
        "packaging": 350
    },

    "C": {
        "units": 900,
        "electricity": 2100,
        "fuel": 350,
        "raw_transport": 500,
        "product_transport": 450,
        "waste": 90,
        "packaging": 280
    },

    "D": {
        "units": 1500,
        "electricity": 3800,
        "fuel": 520,
        "raw_transport": 850,
        "product_transport": 800,
        "waste": 210,
        "packaging": 420
    },

    "E": {
        "units": 1100,
        "electricity": 2700,
        "fuel": 390,
        "raw_transport": 620,
        "product_transport": 550,
        "waste": 130,
        "packaging": 310
    }
}


# ============================================================
# 3. REUSABLE FUNCTION
# ============================================================

def calculate_emissions(batch):

    emissions = {
        "Electricity": batch["electricity"] * ELECTRICITY_FACTOR,

        "Fuel/Gas": batch["fuel"] * FUEL_FACTOR,

        "Raw Material Transport":
            batch["raw_transport"] * RAW_TRANSPORT_FACTOR,

        "Product Transport":
            batch["product_transport"] * PRODUCT_TRANSPORT_FACTOR,

        "Waste":
            batch["waste"] * WASTE_FACTOR,

        "Packaging":
            batch["packaging"] * PACKAGING_FACTOR
    }

    return emissions


# ============================================================
# 4. IMPACT CATEGORY
# ============================================================

def get_impact_category(carbon_per_unit):

    if carbon_per_unit < 5:
        return "Low Impact"

    elif carbon_per_unit < 10:
        return "Moderate Impact"

    else:
        return "High Impact"


# ============================================================
# 5. PROCESS ALL BATCHES
# ============================================================

results = {}

for batch_name, batch in batches.items():

    emissions = calculate_emissions(batch)

    total_carbon = sum(emissions.values())

    carbon_per_unit = total_carbon / batch["units"]

    highest_source = max(
        emissions,
        key=emissions.get
    )

    impact_category = get_impact_category(
        carbon_per_unit
    )

    results[batch_name] = {
        "Total Carbon": total_carbon,
        "Carbon Per Unit": carbon_per_unit,
        "Highest Source": highest_source,
        "Impact Category": impact_category,
        "Waste": batch["waste"]
    }


# ============================================================
# 6. DISPLAY INDIVIDUAL BATCH RESULTS
# ============================================================

print("\n" + "=" * 75)
print("        PRODUCTION PLANT CARBON FOOTPRINT")
print("        ROUND 2 - MULTIPLE PRODUCTION BATCHES")
print("=" * 75)

table = []

for batch, result in results.items():

    table.append([
        batch,
        f"{result['Total Carbon']:.2f}",
        f"{result['Carbon Per Unit']:.2f}",
        result["Highest Source"],
        result["Impact Category"]
    ])


print(
    tabulate(
        table,
        headers=[
            "Batch",
            "Total Carbon (kg CO₂)",
            "Carbon / Unit",
            "Highest Source",
            "Impact Category"
        ],
        tablefmt="grid"
    )
)


# ============================================================
# 7. FIND HIGHEST-CARBON BATCH
# ============================================================

highest_carbon_batch = max(
    results,
    key=lambda batch: results[batch]["Total Carbon"]
)


# ============================================================
# 8. FIND LOWEST-CARBON BATCH
# ============================================================

lowest_carbon_batch = min(
    results,
    key=lambda batch: results[batch]["Total Carbon"]
)


# ============================================================
# 9. FIND MOST EFFICIENT BATCH
# ============================================================

most_efficient_batch = min(
    results,
    key=lambda batch: results[batch]["Carbon Per Unit"]
)


# ============================================================
# 10. FIND HIGHEST WASTE BATCH
# ============================================================

highest_waste_batch = max(
    results,
    key=lambda batch: results[batch]["Waste"]
)


# ============================================================
# 11. DATA AGGREGATION
# ============================================================

total_plant_emissions = sum(
    result["Total Carbon"]
    for result in results.values()
)

average_carbon = (
    total_plant_emissions / len(results)
)


# ============================================================
# 12. NUMPY ANALYSIS
# ============================================================

carbon_values = np.array([
    result["Total Carbon"]
    for result in results.values()
])

carbon_per_unit_values = np.array([
    result["Carbon Per Unit"]
    for result in results.values()
])

print("\nNUMPY ANALYSIS")
print("-" * 75)

print(
    f"Average Carbon : {np.mean(carbon_values):.2f} kg CO₂"
)

print(
    f"Highest Carbon : {np.max(carbon_values):.2f} kg CO₂"
)

print(
    f"Lowest Carbon  : {np.min(carbon_values):.2f} kg CO₂"
)


# ============================================================
# 13. SORTED CONTAINER
# ============================================================

sorted_results = SortedDict()

for batch, result in results.items():
    sorted_results[batch] = result["Total Carbon"]


# ============================================================
# 14. FINAL ANALYSIS
# ============================================================

print("\n" + "=" * 75)
print("DATASET ANALYSIS")
print("=" * 75)

print(
    f"Highest-Carbon Batch : {highest_carbon_batch}"
)

print(
    f"Highest Carbon       : "
    f"{results[highest_carbon_batch]['Total Carbon']:.2f} kg CO₂"
)

print(
    f"\nLowest-Carbon Batch  : {lowest_carbon_batch}"
)

print(
    f"Lowest Carbon        : "
    f"{results[lowest_carbon_batch]['Total Carbon']:.2f} kg CO₂"
)

print(
    f"\nAverage Carbon       : "
    f"{average_carbon:.2f} kg CO₂"
)

print(
    f"\nTotal Plant Emissions: "
    f"{total_plant_emissions:.2f} kg CO₂"
)

print(
    f"\nMost Efficient Batch : "
    f"{most_efficient_batch}"
)

print(
    f"Carbon Per Unit      : "
    f"{results[most_efficient_batch]['Carbon Per Unit']:.2f} kg CO₂"
)

print(
    f"\nHighest Waste Batch  : "
    f"{highest_waste_batch}"
)

print(
    f"Waste Produced       : "
    f"{results[highest_waste_batch]['Waste']} kg"
)


# ============================================================
# 15. PANDAS DATAFRAME
# ============================================================

df = pd.DataFrame.from_dict(
    results,
    orient="index"
)

df.index.name = "Batch"

print("\n" + "=" * 75)
print("PANDAS DATAFRAME")
print("=" * 75)

print(df)


# ============================================================
# 16. PYARROW
# ============================================================

arrow_table = pa.Table.from_pandas(df)

print("\nPyArrow Dataset Created")
print(
    f"Rows: {arrow_table.num_rows}, "
    f"Columns: {arrow_table.num_columns}"
)


# ============================================================
# 17. EXCEL REPORT
# ============================================================

excel_file = "round2_carbon_footprint.xlsx"

with pd.ExcelWriter(
    excel_file,
    engine="xlsxwriter"
) as writer:

    df.to_excel(
        writer,
        sheet_name="Batch Analysis"
    )

    workbook = writer.book
    worksheet = writer.sheets["Batch Analysis"]

    header_format = workbook.add_format({
        "bold": True,
        "border": 1
    })

    number_format = workbook.add_format({
        "num_format": "0.00"
    })

    # Format headers
    for col_num, value in enumerate(
        ["Batch"] + list(df.columns)
    ):
        worksheet.write(
            0,
            col_num,
            value,
            header_format
        )

    worksheet.set_column("A:A", 12)
    worksheet.set_column("B:C", 20)
    worksheet.set_column("D:D", 28)
    worksheet.set_column("E:E", 20)
    worksheet.set_column("F:F", 12)

    # Summary
    worksheet.write("H2", "Dataset Summary", header_format)

    worksheet.write(
        "H3",
        "Highest-Carbon Batch"
    )

    worksheet.write(
        "I3",
        highest_carbon_batch
    )

    worksheet.write(
        "H4",
        "Lowest-Carbon Batch"
    )

    worksheet.write(
        "I4",
        lowest_carbon_batch
    )

    worksheet.write(
        "H5",
        "Average Carbon"
    )

    worksheet.write(
        "I5",
        average_carbon
    )

    worksheet.write(
        "H6",
        "Total Plant Emissions"
    )

    worksheet.write(
        "I6",
        total_plant_emissions
    )

    worksheet.write(
        "H7",
        "Most Efficient Batch"
    )

    worksheet.write(
        "I7",
        most_efficient_batch
    )

    worksheet.write(
        "H8",
        "Highest Waste Batch"
    )

    worksheet.write(
        "I8",
        highest_waste_batch
    )


# ============================================================
# 18. OPENPYXL VERIFICATION
# ============================================================

workbook = load_workbook(excel_file)

print("\nExcel Report Created:")
print(f"File: {excel_file}")
print(f"Sheets: {workbook.sheetnames}")


# ============================================================
# 19. MATPLOTLIB
# ============================================================

plt.figure(figsize=(10, 6))

plt.bar(
    list(results.keys()),
    carbon_values
)

plt.title(
    "Carbon Footprint by Production Batch"
)

plt.xlabel("Production Batch")

plt.ylabel(
    "Total Carbon Footprint (kg CO₂)"
)

plt.tight_layout()

plt.savefig(
    "round2_carbon_comparison.png",
    dpi=150
)

plt.show()


# ============================================================
# 20. PROGRAM COMPLETE
# ============================================================

print("\n" + "=" * 75)
print("ROUND 2 COMPLETE")
print("=" * 75)

print(
    f"Total Plant Carbon: "
    f"{total_plant_emissions:.2f} kg CO₂"
)

print(
    f"Average per Batch: "
    f"{average_carbon:.2f} kg CO₂"
)

print(
    f"Most Efficient: Batch {most_efficient_batch}"
)

print(
    f"Highest Waste: Batch {highest_waste_batch}"
)